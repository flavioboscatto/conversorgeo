"""Gera docs/Manual_Conversor_Geo.docx (manual do aluno).

Ferramenta de documentação, fora do app: requer `python-docx` num ambiente próprio.
    python -m pip install python-docx
    python docs/gerar_manual.py
Figuras: marcadores [FIGURA N — ...] para o professor inserir as capturas de tela.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERDE = RGBColor(0x2E, 0x7D, 0x32)
URL = "https://conversorgeo.com.br"
ACESSO = "Acesso em: 28 set. 2026."
SAIDA = Path(__file__).with_name("Manual_Conversor_Geo.docx")

doc = Document()
_tab = [0]
_fig = [0]


# ------------------------------------------------------------------ utilitários
def _campo(paragrafo, instrucao: str, texto: str = ""):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instrucao)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = texto
    r.append(t)
    fld.append(r)
    paragrafo._p.append(fld)


def _sombrear(celula, cor_hex: str):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    tcPr.append(shd)


def _runs(paragrafo, texto: str):
    """Texto com **negrito** e `código` simples."""
    import re
    for parte in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            paragrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`"):
            r = paragrafo.add_run(parte[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        else:
            paragrafo.add_run(parte)


def titulo(texto: str, nivel: int = 1):
    h = doc.add_heading(texto, level=nivel)
    for r in h.runs:
        r.font.color.rgb = VERDE
    return h


def par(texto: str, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = alinhar
    _runs(p, texto)
    return p


def itens(lista: list[str]):
    for t in lista:
        _runs(doc.add_paragraph(style="List Bullet"), t)


def passos(lista: list[str]):
    """Passos numerados à mão (numeração automática do Word reinicia mal entre listas)."""
    for i, t in enumerate(lista, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{i}. ").bold = True
        _runs(p, t)


def figura(descricao: str):
    _fig[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"[FIGURA {_fig[0]} — {descricao}]")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
    leg = doc.add_paragraph()
    leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    leg.add_run(f"Figura {_fig[0]} — {descricao.split(':')[0]}").italic = True
    leg.add_run("\nFonte: o autor (Conversor Geo).").font.size = Pt(9)


def tabela(legenda: str, cabecalho: list[str], linhas: list[list[str]], larguras_cm: list[float],
           fonte: str = "o autor."):
    _tab[0] += 1
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Tabela {_tab[0]} — {legenda}").bold = True
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, texto in enumerate(cabecalho):
        c = t.rows[0].cells[j]
        c.text = ""
        c.paragraphs[0].add_run(texto).bold = True
        _sombrear(c, "E8F5E9")
    for linha in linhas:
        cells = t.add_row().cells
        for j, texto in enumerate(linha):
            cells[j].text = ""
            _runs(cells[j].paragraphs[0], texto)
    for row in t.rows:
        for j, w in enumerate(larguras_cm):
            row.cells[j].width = Cm(w)
            for p in row.cells[j].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    f = doc.add_paragraph()
    f.add_run(f"Fonte: {fonte}").font.size = Pt(9)


def nota(titulo_nota: str, texto: str, cor: str = "FFF8E1"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.width = Cm(16)
    _sombrear(c, cor)
    c.text = ""
    p = c.paragraphs[0]
    p.add_run(titulo_nota + " ").bold = True
    _runs(p, texto)
    doc.add_paragraph()


def quebra():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ estilo e página
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.orientation = WD_ORIENT.PORTRAIT
sec.top_margin = sec.left_margin = Cm(3)
sec.bottom_margin = sec.right_margin = Cm(2)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
rodape = sec.footer.paragraphs[0]
rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rodape.add_run("Conversor Geo — Manual do Aluno · página ").font.size = Pt(9)
_campo(rodape, "PAGE", "1")
upd = OxmlElement("w:updateFields")          # Word atualiza o sumário ao abrir
upd.set(qn("w:val"), "true")
doc.settings.element.append(upd)

# ------------------------------------------------------------------ capa
for _ in range(6):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CONVERSOR GEO")
r.bold, r.font.size, r.font.color.rgb = True, Pt(30), VERDE
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Manual do Aluno").font.size = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Conversão entre DXF, KML/KMZ, Shapefile e planilha SIGEF, "
          "com visualização sobre imagem e croqui de localização").italic = True
for _ in range(8):
    doc.add_paragraph()
for linha in ("Curso Técnico em Agrimensura", "Instituto Federal de Santa Catarina — IFSC",
              f"Endereço do aplicativo: {URL}", "Versão 1.0 — setembro de 2026"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(linha)
quebra()

titulo("Sumário", 1)
_campo(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
       "Clique com o botão direito aqui e escolha “Atualizar campo” para montar o sumário.")
quebra()

# ------------------------------------------------------------------ 1 objetivos
titulo("1 Objetivos de aprendizagem", 1)
par("Ao final deste material, o aluno deve ser capaz de:")
itens([
    "**configurar** o sistema de coordenadas de entrada (geodésicas ou UTM, com fuso e "
    "hemisfério) de acordo com a origem de cada arquivo;",
    "**converter** levantamentos entre DXF, KML/KMZ e Shapefile, inclusive a partir da "
    "planilha do SIGEF;",
    "**conferir** o resultado sobre imagem de satélite ou ortofoto antes de usá-lo;",
    "**identificar** as mensagens de aviso do aplicativo e corrigir a causa de cada uma;",
    "**produzir** um croqui de localização em imagem (PNG ou JPG) com norte, escala e rótulos.",
])

# ------------------------------------------------------------------ 2 contexto
titulo("2 Contexto prático", 1)
par("No escritório de agrimensura, o mesmo levantamento circula por programas diferentes: o "
    "desenho técnico é feito em CAD (DXF), o cliente confere a área no Google Earth (KML/KMZ), "
    "o órgão ambiental ou a prefeitura pede Shapefile, e a certificação no INCRA parte da "
    "planilha de dados do SIGEF. Cada conversão feita à mão é uma oportunidade de erro de "
    "fuso, de hemisfério ou de sistema de referência — erros que deslocam o imóvel de "
    "centenas de metros a centenas de quilômetros.")
par("O Conversor Geo reúne essas conversões num só lugar, no navegador, sem instalação. "
    "A ideia central é **conferir antes de entregar**: o arquivo é desenhado sobre a imagem, "
    "e só depois convertido.")

# ------------------------------------------------------------------ 3 fundamentação
titulo("3 Fundamentação", 1)
titulo("3.1 Sistema de referência: SIRGAS 2000", 2)
par("O SIRGAS2000 é o sistema geodésico de referência oficial do Brasil; o período de "
    "transição para sua adoção terminou em 25 de fevereiro de 2015 (IBGE, 2015, Art. 1º). "
    "Todo arquivo lido pelo Conversor Geo é levado internamente a coordenadas geodésicas "
    "SIRGAS 2000 (latitude e longitude em graus decimais, com precisão de 64 bits) e "
    "convertido para o sistema de saída só no momento da gravação.")
par("O formato KML aceita apenas longitude e latitude no WGS 84 (OGC, 2008). Para "
    "posicionamentos realizados após 01/01/1994, o IBGE considera WGS 84 e SIRGAS2000 "
    "“idênticos para fins práticos, não existindo parâmetros de transformação entre eles” "
    "(IBGE, 2015, item 2.2). Por isso o aplicativo lê e grava KML sem transformação de datum.")

titulo("3.2 Coordenadas geodésicas e UTM", 2)
par("Na projeção UTM, cada fuso tem 6° de amplitude em longitude, e o meridiano central (MC) "
    "de um fuso é obtido por:")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("MC = 6 × fuso − 183").bold = True
par("em que MC é dado em graus (negativo a oeste de Greenwich). No hemisfério Sul, as "
    "coordenadas N recebem a constante de 10 000 000 m (IBGE, 1999). O aplicativo exige que "
    "fuso e hemisfério sejam **informados** sempre que a entrada ou a saída for UTM: o mesmo "
    "par E/N representa lugares diferentes em fusos diferentes. Ele também sugere o fuso "
    "pela longitude e avisa quando há pontos a mais de 3° do meridiano central do fuso "
    "escolhido.")
tabela("Códigos EPSG usados pelo aplicativo (SIRGAS 2000)",
       ["Sistema", "Código EPSG", "Exemplo"],
       [["Geodésicas (latitude/longitude)", "4674", "—"],
        ["UTM Sul, fusos 17 a 25", "31960 + fuso", "22S → 31982; 23S → 31983"],
        ["UTM Norte, fusos 17 a 22", "31954 + fuso", "22N → 31976"],
        ["UTM Norte, fusos 23 e 24", "6210 e 6211", "31954 + 23 = 31977 seria o 17S!"]],
       [6.0, 3.5, 6.5],
       "IOGP (2026); códigos conferidos por testes automáticos do aplicativo.")
par("Esses códigos aparecem no arquivo `.prj` do Shapefile gerado e são úteis para "
    "configurar o SRC (sistema de referência de coordenadas) no QGIS.")

titulo("3.3 Formatos aceitos", 2)
tabela("Entradas e saídas do Conversor Geo",
       ["Formato", "Entrada", "Saída", "Observações"],
       [["DXF", "Sim", "Sim",
         "Layers, cores e espessuras preservadas. O sistema é informado pelo usuário."],
        ["KML / KMZ", "Sim", "Sim",
         "Sempre geodésico (WGS 84). Uma pasta por layer; KMZ leva o ícone de ponto embutido."],
        ["Shapefile (.zip)", "Sim", "Sim",
         "Entrada: ZIP com .shp, .shx, .dbf e, de preferência, .prj. Saída: ZIP com "
         "_pontos, _linhas e _poligonos, .prj e .cpg (UTF-8)."],
        ["Planilha SIGEF (.ods)", "Sim", "Não",
         "Lê só as abas perimetro_N. A aba de identificação (nome, CPF) não é lida."],
        ["DWG", "Não", "Não", "Salvar como DXF no programa de CAD antes de usar."]],
       [3.2, 1.7, 1.6, 9.5])

# ------------------------------------------------------------------ 4 passo a passo
quebra()
titulo("4 Passo a passo", 1)
titulo("4.1 Acesso", 2)
passos([
    f"Abrir o endereço **{URL}** no navegador (Chrome, Edge ou Firefox atualizados).",
    "Se a página demorar cerca de 1 minuto no primeiro acesso, aguardar: o servidor "
    "gratuito pode estar “acordando”.",
])
nota("Privacidade:", "os arquivos enviados ficam só na memória da sessão e não são gravados "
     "no servidor. Ao fechar a página, eles são descartados.", "E3F2FD")
figura("tela inicial do Conversor Geo: quadros Entrada, Saída, mapa, Imagem e Layers")

titulo("4.2 Importar um arquivo (Entrada)", 2)
par("A tela segue a ordem **formato → sistema → arquivo**. Cada arquivo tem a sua própria "
    "configuração.")
passos([
    "Em **Formato**, escolher o tipo do arquivo: DXF, KML / KMZ, Shapefile (.zip) ou "
    "Planilha SIGEF (.ods).",
    "Em **Sistema de coordenadas**, escolher **Geodésicas (graus decimais)** ou **UTM**. "
    "Para UTM, informar o **Fuso** e o hemisfério (**Sul** ou **Norte**). Para KML o sistema "
    "fica travado em geodésicas; para a planilha SIGEF o sistema é lido da própria planilha "
    "(campo “Tipo de Coordenada”); para Shapefile vale o .prj, e o sistema escolhido só é "
    "usado se o ZIP não tiver .prj.",
    "Arrastar o arquivo para a área de envio (ou clicar no **+**). Limite: 10 MB por arquivo. Aparece "
    "a mensagem “Arquivo recebido”.",
    "Clicar em **Mostrar arquivo na imagem**. O arquivo é lido, desenhado sobre a imagem e "
    "entra na lista **Arquivos importados**, com o formato, o sistema e a contagem de pontos, "
    "linhas, polígonos e textos.",
])
figura("quadro Entrada com formato DXF, sistema UTM fuso 22 Sul e arquivo recebido")

titulo("4.3 Conferir no mapa", 2)
itens([
    "O mapa se ajusta ao arquivo. Conferir se o desenho caiu **no lugar certo** da imagem — "
    "esse é o principal controle de qualidade da conversão.",
    "No quadro **Layers**, cada layer tem uma caixa para ligar/desligar e um quadrado com a "
    "cor. Clicar no quadrado permite **trocar a cor**; a nova cor vale para o mapa, para a "
    "imagem exportada e para os arquivos de saída.",
    "Clicar numa feição abre os atributos (por exemplo, sigmas, método de posicionamento, "
    "tipo de limite e confrontante dos vértices do SIGEF).",
    "Avisos aparecem em laranja abaixo dos quadros; erros, em vermelho (ver seção 6).",
])
figura("mapa com a parcela do SIGEF sobre a imagem, painel de layers e janela de atributos")

titulo("4.4 Vários arquivos numa saída única", 2)
par("Repetir a seção 4.2 para cada arquivo — cada um com o seu formato e o seu sistema (um "
    "DXF em UTM e uma planilha SIGEF em geodésicas, por exemplo). Todos aparecem juntos no "
    "mapa, e a saída é **sempre um arquivo único** com tudo. Quando dois arquivos têm uma "
    "layer com o mesmo nome, a do segundo recebe o nome do arquivo na frente (ex.: "
    "`lote2_DIVISA`).")
itens([
    "O **✕** de cada item da lista remove só aquele arquivo.",
    "**Limpar importação** remove todos os arquivos.",
])

titulo("4.5 Converter e baixar (Saída)", 2)
passos([
    "Em **Saída — arquivo único**, escolher o **Formato**: DXF, KML / KMZ ou Shapefile (.zip).",
    "Escolher o **Sistema de coordenadas** da saída. O fuso UTM já vem sugerido pelo arquivo "
    "importado; conferir. KML fica travado em geodésicas.",
    "Para KML / KMZ, escolher **KMZ (com ícones)** — recomendado — ou **KML**.",
    "Clicar em **Converter e baixar**. O arquivo é salvo na pasta de downloads do navegador.",
])
tabela("O que cada saída contém",
       ["Saída", "Conteúdo"],
       [["DXF (R2018)", "Uma layer por layer de origem, com cor e espessura. Pontos como POINT "
                        "com rótulo em TEXT ao lado; linhas e polígonos como polilinhas "
                        "(3D quando há altitude variável)."],
        ["KMZ / KML", "Uma pasta por layer; pontos com ícone de círculo colorido e rótulo; "
                      "textos soltos só com rótulo; atributos na janela de cada feição."],
        ["Shapefile (.zip)", "Até três shapefiles (`_pontos`, `_linhas`, `_poligonos`), "
                             "campos LAYER, TIPO, COR, ESPESSURA, TEXTO e atributos de origem "
                             "(nomes com até 10 caracteres), `.prj` e `.cpg`."]],
       [3.5, 12.5])

titulo("4.6 Croqui de localização (imagem)", 2)
passos([
    "No quadro **Imagem**, escolher a **Imagem de fundo**: **Esri (atual)** — padrão, "
    "cobertura mundial — ou **Ortofoto SC 2012 (só SC)**, do aerolevantamento do Estado de "
    "Santa Catarina, com 0,39 m de resolução (SANTA CATARINA, 2026).",
    "Escolher os **Rótulos**: **Sem rótulos**; **Esri (ruas e lugares)**; ou "
    "**OpenStreetMap (só nomes de ruas)**, que escreve os nomes ao longo das ruas, sem "
    "números de endereço, a partir do zoom 15.",
    "No quadro **Layers**, deixar ligadas só as layers que devem aparecer no croqui.",
    "Enquadrar o mapa (zoom e posição) como o croqui deve ficar.",
    "Escolher **PNG** ou **JPG** e clicar em **Exportar imagem**.",
])
par("A imagem sai com o dobro da resolução da tela, seta de norte no canto superior esquerdo, "
    "escala em barra no canto inferior direito e o crédito das fontes usadas (Esri, SDE/SC, "
    "OpenStreetMap), que deve ser mantido no croqui.")
figura("croqui exportado: ortofoto SC, parcela, nomes de ruas do OpenStreetMap, norte e escala")

# ------------------------------------------------------------------ 5 exemplos
quebra()
titulo("5 Exemplos resolvidos", 1)
titulo("5.1 Planilha SIGEF para KMZ", 2)
par("Arquivo: planilha de dados do SIGEF com uma aba `perimetro_1`, “Tipo de Coordenada” "
    "geográfica e 16 vértices (FBSC-M-0001 a FBSC-P-0011).")
passos([
    "Entrada: formato **Planilha SIGEF (.ods)** (o sistema é lido da planilha). Enviar o "
    "arquivo e clicar em **Mostrar arquivo na imagem**.",
    "Conferir o resumo: **16 pontos · 16 linhas · 1 polígono · 0 textos**. As layers "
    "criadas são VERTICES_M, VERTICES_P, VERTICES_V (por tipo de vértice), LIMITES (um "
    "segmento por lado, com tipo de limite, CNS, matrícula e confrontante) e PERIMETRO.",
    "Conferir a posição do vértice FBSC-M-0001, informado na planilha como "
    "27°35'39,554\" S e 48°32'28,781\" W.",
    "Saída: **KML / KMZ**, **KMZ (com ícones)**, **Converter e baixar**. Abrir o KMZ no "
    "Google Earth.",
])
par("Verificação do fuso para uma saída em UTM: a longitude do vértice é −48,5413°, que cai "
    "no fuso 22 (de −54° a −48°, MC = 6 × 22 − 183 = −51°). No fuso 22 Sul, o aplicativo "
    "calcula para o FBSC-M-0001 **E = 742 684,931 m** e **N = 6 945 322,592 m**.")

titulo("5.2 DXF em UTM para Shapefile", 2)
par("Arquivo: `exemplo_utm22s.dxf` (fornecido pelo professor), desenhado em UTM 22 Sul, com "
    "vértices P1 a P4, divisa com trecho em arco, edificações, arruamento, um bloco de marco "
    "e um objeto esquecido na origem (0, 0) do desenho.")
passos([
    "Entrada: **DXF**, **UTM**, fuso **22**, **Sul**. Enviar e mostrar na imagem.",
    "Resumo esperado: **5 pontos · 3 linhas · 3 polígonos · 1 texto**. Os textos P1 a P4 "
    "estão a menos de 0,5 m dos pontos e viram **rótulos** dos pontos; o texto “Área de "
    "teste” fica solto.",
    "Aviso esperado: “1 feição(ões) com coordenadas fora da faixa de SIRGAS 2000 / UTM 22S "
    "foram ignoradas” — é o objeto na origem, que não pertence ao levantamento.",
    "Saída: **Shapefile (.zip)**, **UTM** fuso **22 Sul**, **Converter e baixar**. O ZIP "
    "traz `exemplo_utm22s_pontos`, `_linhas` e `_poligonos`, cada um com .shp, .shx, .dbf, "
    ".prj e .cpg.",
    "No QGIS, arrastar o ZIP para o mapa: o SRC é reconhecido pelo .prj como "
    "SIRGAS 2000 / UTM zone 22S (EPSG:31982).",
])

# ------------------------------------------------------------------ 6 erros
quebra()
titulo("6 Erros comuns e cuidados", 1)
tabela("Mensagens do aplicativo: causa e solução",
       ["Mensagem (resumo)", "Causa provável", "O que fazer"],
       [["“As coordenadas parecem UTM…”", "Arquivo em metros, sistema escolhido: Geodésicas.",
         "Escolher UTM, com fuso e hemisfério corretos."],
        ["“As coordenadas parecem geodésicas (graus)…”", "Arquivo em graus, sistema escolhido: UTM.",
         "Escolher Geodésicas."],
        ["“…nenhuma coordenada cabe em … reprojete…”",
         "Arquivo em outro sistema (ex.: Web Mercator, comum em exportações do QGIS).",
         "No QGIS: Exportar → Salvar feições como, escolhendo SIRGAS 2000 / UTM ou geodésicas; "
         "importar de novo."],
        ["“… a mais de 3° do meridiano central do fuso …”",
         "Fuso informado diferente do fuso real dos pontos.",
         "Usar o fuso sugerido pela longitude."],
        ["“… no hemisfério Norte/Sul, mas o sistema informado é …”",
         "Hemisfério trocado.", "Corrigir Sul/Norte."],
        ["“… feição(ões) … fora da faixa … ignoradas”",
         "Objetos soltos longe do levantamento (ex.: na origem 0,0 do CAD).",
         "Conferir se não falta nada no mapa; apagar esses objetos no CAD, se for o caso."],
        ["“… não tem .prj; usado o sistema informado”", "Shapefile sem arquivo .prj.",
         "Conferir se o sistema escolhido na entrada é o do arquivo."],
        ["“Falta o arquivo ….dbf dentro do ZIP”", "ZIP incompleto.",
         "Compactar juntos .shp, .shx, .dbf e .prj."],
        ["“O arquivo não parece ser um DXF válido”", "Arquivo DWG renomeado ou DXF corrompido.",
         "No CAD: Salvar como → DXF."],
        ["“Servidor do OpenStreetMap indisponível…”", "Serviço público de dados ocupado.",
         "Mover o mapa para tentar de novo, ou usar os rótulos Esri."]],
       [4.6, 5.0, 6.4])
titulo("Cuidados", 2)
itens([
    "**Linhas brancas do AutoCAD** (cor 7) passam a **pretas** nas saídas, porque é assim que "
    "Google Earth, QGIS e impressão esperam. No DXF de saída, preto volta como cor 7. Para "
    "outra cor, trocar no quadro Layers.",
    "**Tracejados** não existem no KML: a linha sai contínua.",
    "**Arcos e círculos** são convertidos em segmentos de 1° em 1°.",
    "**Furos de polígonos** (anéis internos) não são convertidos.",
    "**Rótulos Esri** mostram números de endereço a partir do zoom 18 (escala de 50 m); para "
    "croquis mais aproximados sem números, usar os rótulos do OpenStreetMap.",
    "**Nomes de ruas** podem divergir entre Esri e OpenStreetMap; conferir com a prefeitura "
    "quando o nome for importante para o trabalho.",
    "A **Ortofoto SC** é de 2010–2012 e cobre só Santa Catarina; fora do estado aparece a "
    "imagem Esri. É uma ortofoto verdadeira (sem inclinação das edificações), mas "
    "desatualizada — construções recentes não aparecem.",
    "A imagem de fundo serve para **conferência visual e croqui**, não para medição.",
])

# ------------------------------------------------------------------ 7 exercícios
quebra()
titulo("7 Exercícios de fixação", 1)
passos([
    "Um vértice tem longitude 48°32'28,781\" W. Em qual fuso UTM ele está e qual é o "
    "meridiano central desse fuso?",
    "Qual é o código EPSG de SIRGAS 2000 / UTM 23 Sul? E de UTM 24 Norte?",
    "Um colega recebeu um KML do Google Earth e quer convertê-lo para DXF em UTM. Qual "
    "sistema deve escolher na **entrada**? Por que não há transformação de datum entre o "
    "KML e o SIRGAS 2000?",
    "Importar `exemplo_utm22s.dxf` como **Geodésicas**. Qual mensagem aparece e por quê? "
    "Corrigir e registrar o resumo de feições obtido.",
    "No exemplo 5.2, por que os textos P1 a P4 não aparecem como “textos” no resumo?",
    "Importar a planilha SIGEF e o `exemplo_utm22s.dxf` juntos e gerar **um** KMZ. Quantas "
    "pastas (layers) o KMZ terá?",
    "Produzir um croqui de localização da parcela do SIGEF em JPG, com Ortofoto SC, rótulos "
    "do OpenStreetMap e só as layers PERIMETRO e VERTICES_M ligadas. Em que zoom os rótulos "
    "Esri passariam a mostrar números de endereço?",
    "Desafio: um DXF exportado do QGIS tem vértices com X ≈ −5 401 872 e Y ≈ −3 198 591. "
    "O aplicativo não o importa em UTM nem em geodésicas. Explicar a causa provável e o "
    "procedimento de correção.",
])

# ------------------------------------------------------------------ gabarito
quebra()
titulo("Gabarito", 1)
passos([
    "Longitude −48,5413°: fuso 22 (de −54° a −48°); MC = 6 × 22 − 183 = **−51°**.",
    "UTM 23 Sul: **31983** (31960 + 23). UTM 24 Norte: **6211** (a regra 31954 + fuso só "
    "vale até o fuso 22 Norte).",
    "**Geodésicas** — o sistema já fica travado para KML. O KML usa WGS 84, e para "
    "posicionamentos após 01/01/1994 o IBGE considera WGS 84 e SIRGAS2000 idênticos para "
    "fins práticos (IBGE, 2015, item 2.2).",
    "“As coordenadas do DXF parecem UTM (valores em metros…)”: os números estão em metros, "
    "incompatíveis com graus. Com UTM 22 Sul: 5 pontos · 3 linhas · 3 polígonos · 1 texto.",
    "Porque estão a menos de 0,5 m de um ponto e são associados a ele como **rótulo**; só o "
    "texto “Área de teste”, afastado dos pontos, permanece como texto solto.",
    "**10 pastas**: 5 layers da planilha (VERTICES_M, VERTICES_P, VERTICES_V, LIMITES, "
    "PERIMETRO) + 5 do DXF (ARRUAMENTO, DIVISA, EDIFICACAO, TEXTOS, VERTICES). Não há nomes "
    "repetidos, portanto nenhum prefixo é acrescentado.",
    "Resposta prática (conferir a imagem). Os números de endereço aparecem a partir do "
    "**zoom 18** (barra de escala de 50 m).",
    "Os valores são de **Web Mercator (EPSG:3857)**, padrão de mapas web e de muitos projetos "
    "do QGIS, que o aplicativo não aceita. Reprojetar no QGIS (Exportar → Salvar feições "
    "como, com SRC SIRGAS 2000 / UTM zone 22S ou SIRGAS 2000 geodésico) e importar de novo.",
])

# ------------------------------------------------------------------ referências
titulo("Referências", 1)
refs = [
    "INSTITUTO BRASILEIRO DE GEOGRAFIA E ESTATÍSTICA (IBGE). **Noções básicas de cartografia**. "
    "Rio de Janeiro: IBGE, 1999. (Manuais técnicos em geociências, n. 8).",
    "INSTITUTO BRASILEIRO DE GEOGRAFIA E ESTATÍSTICA (IBGE). **Resolução R.PR n. 1, de 24 de "
    "fevereiro de 2015**. Define a data de término do período de transição definido na RPR "
    "01/2005 e dá outras providências sobre a transformação entre os referenciais geodésicos "
    "adotados no Brasil. Rio de Janeiro: IBGE, 2015. Disponível em: "
    "https://geoftp.ibge.gov.br/metodos_e_outros_documentos_de_referencia/normas/"
    f"rpr_01_2015_sirgas2000.pdf. {ACESSO}",
    "INTERNATIONAL ASSOCIATION OF OIL & GAS PRODUCERS (IOGP). **EPSG Geodetic Parameter "
    f"Dataset**. [S. l.]: IOGP, 2026. Disponível em: https://epsg.org. {ACESSO}",
    "OPEN GEOSPATIAL CONSORTIUM (OGC). **OGC KML**: version 2.2.0. OGC 07-147r2. "
    "Wayland: OGC, 2008. Disponível em: https://www.ogc.org/standards/kml. " + ACESSO,
    "OPENSTREETMAP. **Copyright and License**. [S. l.]: OpenStreetMap Foundation, 2026. "
    f"Disponível em: https://www.openstreetmap.org/copyright. {ACESSO}",
    "SANTA CATARINA. Secretaria de Estado do Desenvolvimento Econômico Sustentável. "
    "**SIG@SC**: URLs de serviços WMS e WMTS. Florianópolis: SDE, 2026. Disponível em: "
    f"https://sigsc.sc.gov.br/wms-urls.html. {ACESSO}",
]
for ref in refs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(8)
    _runs(p, ref)

doc.save(SAIDA)
print(f"Gerado: {SAIDA}  ({_tab[0]} tabelas, {_fig[0]} figuras marcadas)")
